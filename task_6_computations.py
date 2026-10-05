from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# =========================================================
# CONFIGURATION
# =========================================================

INPUT_FILE = Path("221953930-219431865.xlsx")
OUTPUT_FILE = Path("221953930-219431865_with_stats.xlsx")

STAT_SHEETS = [
    "Stat_Nominal",
    "Stat_Ordinal",
    "Stat_Interval",
    "Stat_Ratio",
]


# =========================================================
# HELPERS
# =========================================================

def get_column_index(ws, header):
    """Return column number for a header in row 1."""
    for cell in ws[1]:
        if cell.value == header:
            return cell.column
    raise ValueError(f"Column '{header}' not found in sheet '{ws.title}'")


def get_column_values(ws, header):
    """Return all non-empty values below a header."""
    col = get_column_index(ws, header)

    return [
        ws.cell(row=row, column=col).value
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=col).value is not None
    ]


def style_sheet(ws):
    """Simple formatting."""
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.freeze_panes = "A2"

    # Reasonable widths
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 22


# =========================================================
# LOAD WORKBOOK
# =========================================================

wb = load_workbook(INPUT_FILE)

stock_profile = wb["Stock_Profile"]
stock_day = wb["Stock_Day"]
intraday = wb["Intraday_Bars"]

# Remove old statistics sheets if rerunning script
for name in STAT_SHEETS:
    if name in wb.sheetnames:
        del wb[name]


# =========================================================
# READ SOURCE DATA
# =========================================================

# Stock_Profile
profile_tickers = get_column_values(stock_profile, "ticker")
exchanges = get_column_values(stock_profile, "exchange")
mean_daily_volumes = get_column_values(stock_profile, "mean_daily_volume")
liquidity_tiers = get_column_values(stock_profile, "liquidity_tier")

# Stock_Day
day_tickers = get_column_values(stock_day, "ticker")
session_dates = get_column_values(stock_day, "session_date")
session_volumes = get_column_values(stock_day, "session_volume")

# Intraday_Bars
bar_tickers = get_column_values(intraday, "ticker")
bar_times = get_column_values(intraday, "bar_start_utc")


# =========================================================
# 1. NOMINAL — EXCHANGE
# =========================================================

ws = wb.create_sheet("Stat_Nominal")

ws.append(["ticker", "exchange"])

for ticker, exchange in zip(profile_tickers, exchanges):
    ws.append([ticker, exchange])

last_row = ws.max_row

# Summary area
ws["D1"] = "Statistic"
ws["E1"] = "Result"
ws["D2"] = "Number of observations"
ws["E2"] = f"=COUNTA(B2:B{last_row})"

ws["D4"] = "Exchange"
ws["E4"] = "Frequency"
ws["F4"] = "Proportion"

unique_exchanges = sorted(set(exchanges))

for row, exchange in enumerate(unique_exchanges, start=5):
    ws.cell(row=row, column=4, value=exchange)
    ws.cell(
        row=row,
        column=5,
        value=f'=COUNTIF($B$2:$B${last_row},D{row})'
    )
    ws.cell(
        row=row,
        column=6,
        value=f"=E{row}/$E$2"
    )
    ws.cell(row=row, column=6).number_format = "0.00%"

freq_end = 4 + len(unique_exchanges)

ws["D3"] = "Mode"
ws["E3"] = (
    f"=INDEX(D5:D{freq_end},"
    f"MATCH(MAX(E5:E{freq_end}),E5:E{freq_end},0))"
)

style_sheet(ws)


# =========================================================
# 2. ORDINAL — LIQUIDITY TIER
# =========================================================

ws = wb.create_sheet("Stat_Ordinal")

ws.append([
    "ticker",
    "mean_daily_volume",
    "liquidity_tier",
    "tier_rank"
])

for ticker, volume, tier in zip(
    profile_tickers,
    mean_daily_volumes,
    liquidity_tiers
):
    ws.append([ticker, volume, tier])

last_row = ws.max_row

# Excel helper ranking
for row in range(2, last_row + 1):
    ws.cell(
        row=row,
        column=4,
        value=f'=IF(C{row}="Low",1,IF(C{row}="Medium",2,IF(C{row}="High",3,"")))'
    )

# Summary
ws["F1"] = "Statistic"
ws["G1"] = "Result"

ws["F2"] = "Number of observations"
ws["G2"] = f"=COUNTA(C2:C{last_row})"

ws["F3"] = "Median category"
ws["G3"] = (
    f'=CHOOSE(MEDIAN(D2:D{last_row}),'
    '"Low","Medium","High")'
)

ws["F4"] = "Mode(s)"

# Frequency table
ws["F7"] = "Liquidity Tier"
ws["G7"] = "Frequency"
ws["H7"] = "Proportion"

tiers = ["Low", "Medium", "High"]

for row, tier in enumerate(tiers, start=8):
    ws.cell(row=row, column=6, value=tier)
    ws.cell(
        row=row,
        column=7,
        value=f'=COUNTIF($C$2:$C${last_row},F{row})'
    )
    ws.cell(
        row=row,
        column=8,
        value=f"=G{row}/$G$2"
    )
    ws.cell(row=row, column=8).number_format = "0.00%"

# Returns tied modes too
ws["G4"] = (
    '=TEXTJOIN(", ",TRUE,'
    'IF(G8=MAX($G$8:$G$10),F8,""),'
    'IF(G9=MAX($G$8:$G$10),F9,""),'
    'IF(G10=MAX($G$8:$G$10),F10,""))'
)

for row in range(2, last_row + 1):
    ws.cell(row=row, column=2).number_format = "#,##0.00"

style_sheet(ws)


# =========================================================
# 3. INTERVAL — BAR_START_UTC
# =========================================================

ws = wb.create_sheet("Stat_Interval")

ws.append(["ticker", "bar_start_utc"])

for ticker, timestamp in zip(bar_tickers, bar_times):

    # If timestamp was stored as a string, try converting it
    if isinstance(timestamp, str):
        text = timestamp.replace("Z", "+00:00")

        try:
            timestamp = datetime.fromisoformat(text)

            # Excel cannot store timezone-aware datetime objects
            if timestamp.tzinfo is not None:
                timestamp = timestamp.replace(tzinfo=None)

        except ValueError:
            pass

    ws.append([ticker, timestamp])

last_row = ws.max_row

# Date formatting
for row in range(2, last_row + 1):
    ws.cell(row=row, column=2).number_format = "yyyy-mm-dd hh:mm:ss"

# Summary
ws["D1"] = "Statistic"
ws["E1"] = "Result"

ws["D2"] = "Number of observations"
ws["E2"] = f"=COUNT(B2:B{last_row})"

ws["D3"] = "Mean timestamp"
ws["E3"] = f"=AVERAGE(B2:B{last_row})"
ws["E3"].number_format = "yyyy-mm-dd hh:mm:ss"

ws["D4"] = "Sample standard deviation (hours)"
ws["E4"] = f"=STDEV.S(B2:B{last_row})*24"
ws["E4"].number_format = "0.00"

style_sheet(ws)


# =========================================================
# 4. RATIO — SESSION_VOLUME
# =========================================================

ws = wb.create_sheet("Stat_Ratio")

ws.append([
    "ticker",
    "session_date",
    "session_volume"
])

for ticker, date, volume in zip(
    day_tickers,
    session_dates,
    session_volumes
):
    ws.append([ticker, date, volume])

last_row = ws.max_row

for row in range(2, last_row + 1):
    ws.cell(row=row, column=2).number_format = "yyyy-mm-dd"
    ws.cell(row=row, column=3).number_format = "#,##0"

# Required calculations
ws["E1"] = "Statistic"
ws["F1"] = "Result"

ws["E2"] = "Number of observations"
ws["F2"] = f"=COUNT(C2:C{last_row})"

ws["E3"] = "Arithmetic mean"
ws["F3"] = f"=AVERAGE(C2:C{last_row})"

ws["E4"] = "Sample standard deviation"
ws["F4"] = f"=STDEV.S(C2:C{last_row})"

# Optional descriptive statistics
ws["E6"] = "Median"
ws["F6"] = f"=MEDIAN(C2:C{last_row})"

ws["E7"] = "Minimum"
ws["F7"] = f"=MIN(C2:C{last_row})"

ws["E8"] = "Maximum"
ws["F8"] = f"=MAX(C2:C{last_row})"

for cell in ["F3", "F4", "F6", "F7", "F8"]:
    ws[cell].number_format = "#,##0.00"

style_sheet(ws)


# =========================================================
# SAVE
# =========================================================

wb.save(OUTPUT_FILE)

print(f"Done. Saved workbook as: {OUTPUT_FILE}")
print()
print("Created sheets:")
for name in STAT_SHEETS:
    print(" -", name)